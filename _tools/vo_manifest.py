# -*- coding: utf-8 -*-
"""Write VO_Manifest.xlsx - every clip the game plays, with its script, for the recording session.

    python _tools/vo_manifest.py            (after _tools/build_card.py)

Same sheet as the fleet's other manifests (MTG2A01_L06_S01/VO_Manifest.xlsx): one sheet "VO Manifest",
two columns Audio ID | Script, navy header, frozen first row, filter on.

ONLY CLIPS THE GAME PLAYS, IN THE ORDER IT PLAYS THEM. Read from card.json: the landing, each phase
gate where its round begins, then every slide's own clips (prompt, option sentences, sentence, hints,
reveal, correct). The single-word clips spoken when a tile or option is tapped come last, as one block -
easier to record in one go. The TTS building blocks build_card.py joins (vo_sahi, vo_shabash_sahi) are
not played on their own, so they are not listed.

RECORD EACH ROW AS ONE CLIP, saved as <Audio ID>.wav. "…" in a script is a pause where the blank is.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")
PROJ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
CARD = json.load(open(os.path.join(PROJ, "HI02H12_L02_S01", "card.json"), encoding="utf-8"))
OUT = os.path.join(PROJ, "VO_Manifest.xlsx")
GATE = {"tutorial": "vo_pt_tutorial", "guided": "vo_pt_guided", "practice": "vo_pt_practice"}
SLOTS = ["prompt", "sentence", "hint1", "hint2", "reveal", "correct"]


def ids(v): return v if isinstance(v, list) else [v]


def main():
    text = CARD["assets"]["audio_text"]
    order, words, seen = ["vo_landing"], [], {"vo_landing"}
    def add(i, into=order):
        if i and i.startswith("vo_") and i not in seen: seen.add(i); into.append(i)
    phase = None
    for s in CARD["slides"]:
        if s["phase"] != phase and s["type"] != "CELEBRATION":
            phase = s["phase"]; add(GATE.get(phase))
        a, d = s.get("audio", {}), s.get("data", {})
        for i in ids(a.get("prompt")): add(i)
        for o in d.get("options", []): add(o.get("say"))          # I Do: each option's sentence
        for k in SLOTS[1:]:
            for i in ids(a.get(k)): add(i)
        for o in d.get("options", []) + d.get("tiles", []): add(o.get("audio"), words)
    rows = order + words

    wb = Workbook(); ws = wb.active; ws.title = "VO Manifest"
    ws.append(["Audio ID", "Script"])
    for i in rows: ws.append([i, text.get(i, "")])
    head = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, size=12, color="FFFFFF"); c.fill = head
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 118
    ws.freeze_panes = "A2"; ws.auto_filter.ref = "A1:B%d" % ws.max_row
    wb.save(OUT)
    missing = [i for i in rows if not text.get(i)]
    print("VO_Manifest.xlsx:", len(rows), "clips (%d words)" % len(words), "| no script:", missing or "none")


if __name__ == "__main__":
    main()
